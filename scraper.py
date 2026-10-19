"""ACNC charity scraper (portfolio sample).

Searches charities on acnc.gov.au via its internal JSON APIs, keeps only
registered ones, fetches each detail record (with a polite delay) and
exports ABN, Address for Service Email and Charity size to Excel.
Progress is saved after every record, so an interrupted run can resume.
The delay between requests adapts to the server (slows down on 429/403).
Records are processed in batches; the Excel file is saved after each batch.
"""

import json
import os
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://www.acnc.gov.au/api/dynamics"
SEARCH_TERM = "garden"
SAMPLE_SIZE = 100
BATCH_SIZE = 10
DELAY_SECONDS = 1
OUTPUT_FILE = "acnc_charities_sample.xlsx"
PROGRESS_FILE = "progress.json"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json",
}

MAX_DELAY = 30
MAX_RETRIES = 3


class AdaptiveDelay:
    """Pause between requests: grows on 429/403, shrinks back when OK."""

    def __init__(self, base):
        self.base = base
        self.current = base

    def wait(self):
        time.sleep(self.current)

    def slow_down(self):
        self.current = min(self.current * 2, MAX_DELAY)

    def speed_up(self):
        self.current = max(self.current * 0.8, self.base)


throttle = AdaptiveDelay(DELAY_SECONDS)


def http_get(url, **kwargs):
    """GET with retries; backs off when the server answers 429 or 403."""
    for attempt in range(1, MAX_RETRIES + 1):
        resp = requests.get(url, headers=HEADERS, timeout=30, **kwargs)
        if resp.status_code in (429, 403):
            throttle.slow_down()
            print(
                f"  Blocked ({resp.status_code}). Pausing {throttle.current:.0f}s "
                f"(attempt {attempt}/{MAX_RETRIES})"
            )
            throttle.wait()
            continue
        resp.raise_for_status()
        throttle.speed_up()
        return resp
    raise requests.RequestException(f"Still blocked after {MAX_RETRIES} attempts")


def load_progress(path):
    """Return {uuid: row or None} from a previous run (empty if none)."""
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_progress(progress, path):
    """Save progress safely (temp file first, so a crash can't corrupt it)."""
    tmp_path = path + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False)
    os.replace(tmp_path, path)


def search_charities(term, sample_size):
    """Return up to `sample_size` registered charities matching `term`."""
    found = []
    page = 0
    while len(found) < sample_size:
        params = {"search": term, "page": page, "items_per_page": 50}
        resp = http_get(f"{BASE_URL}/search/charity", params=params)
        payload = resp.json()

        for item in payload.get("results", []):
            data = item.get("data", {})
            if data.get("Status") == "Registered":
                found.append({"uuid": item["uuid"], "name": data.get("Name")})
                if len(found) >= sample_size:
                    break

        total_pages = payload.get("pager", {}).get("total_pages", 0)
        page += 1
        if page >= total_pages:
            break
        throttle.wait()
    return found


def fetch_detail(uuid):
    """Return the detail record (dict) for one charity."""
    resp = http_get(f"{BASE_URL}/entity/{uuid}")
    return resp.json().get("data", {})


def export_to_excel(rows, path):
    """Write rows to a formatted Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Charities"

    columns = ["Name", "ABN", "Address for Service Email", "Charity Size"]
    ws.append(columns)
    for row in rows:
        ws.append([row["name"], row["abn"], row["email"], row["size"]])

    header_fill = PatternFill("solid", start_color="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, col in enumerate(columns, start=1):
        longest = max(len(str(c.value or "")) for c in ws[get_column_letter(idx)])
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 3, 60)

    ws.freeze_panes = "A2"
    wb.save(path)


def process_charity(charity, label, progress):
    """Fetch one charity, update progress and save it. Errors are not saved."""
    uuid = charity["uuid"]
    if uuid in progress:
        print(f"{label} DONE (previous run) - {charity['name']}")
        return
    try:
        detail = fetch_detail(uuid)
        if not detail.get("Abn"):
            progress[uuid] = None  # remember it was skipped
            print(f"{label} SKIP (no ABN) - {charity['name']}")
        else:
            progress[uuid] = {
                "name": detail.get("Name") or charity["name"],
                "abn": detail.get("Abn"),
                "email": detail.get("AFSEmail"),
                "size": detail.get("CharitySize") or "Not reported",
            }
            print(f"{label} OK - {charity['name']}")
        save_progress(progress, PROGRESS_FILE)
    except requests.RequestException as exc:
        print(f"{label} ERROR - {charity['name']}: {exc}")
    throttle.wait()


def save_excel(progress):
    """Export all collected rows; warn instead of crashing if the file is open."""
    rows = [row for row in progress.values() if row]
    try:
        export_to_excel(rows, OUTPUT_FILE)
        print(f"  Saved {len(rows)} rows to {OUTPUT_FILE}")
    except PermissionError:
        print(f"  Could not save {OUTPUT_FILE}. Close it in Excel; will retry next batch.")


def main():
    print(f"Searching for '{SEARCH_TERM}'...")
    charities = search_charities(SEARCH_TERM, SAMPLE_SIZE)
    total = len(charities)
    print(f"Found {total} registered charities.")

    progress = load_progress(PROGRESS_FILE)
    print(f"Resuming: {len(progress)} records already processed.")

    total_batches = (total + BATCH_SIZE - 1) // BATCH_SIZE
    for start in range(0, total, BATCH_SIZE):
        print(f"--- Batch {start // BATCH_SIZE + 1}/{total_batches} ---")
        batch = charities[start:start + BATCH_SIZE]
        for i, charity in enumerate(batch, start=start + 1):
            process_charity(charity, f"[{i}/{total}]", progress)
        save_excel(progress)

    print("Done.")


if __name__ == "__main__":
    main()